import csv
from pathlib import Path

import openpyxl
import pytest

from app.models.models import ActivityType, Benchmark, EmissionFactor, Intervention

BASE_DIR = Path(__file__).resolve().parent.parent.parent


def test_seed_idempotency_and_integrity(db_session):
    """Verify scripts/seed.py runs idempotently and populates all canonical tables."""
    import sys

    sys.path.insert(0, str(BASE_DIR / "scripts"))
    from seed import run_seed

    seed_dir = BASE_DIR / "data" / "seed"
    assert seed_dir.exists(), "data/seed directory must exist"

    # Run seed first time
    counts1 = run_seed(session=db_session, seed_dir=seed_dir)
    assert counts1["activity_types"] == 26
    assert counts1["emission_factors"] == 26
    assert counts1["interventions"] == 16
    assert counts1["benchmarks"] == 4

    # Run seed second time (idempotency check)
    counts2 = run_seed(session=db_session, seed_dir=seed_dir)
    assert counts1 == counts2, "Seed must be strictly idempotent"

    # Verify ActivityType content & localization
    grid = db_session.query(ActivityType).filter_by(key="grid_electricity").first()
    assert grid is not None
    assert grid.category == "electricity"
    assert grid.canonical_unit == "kWh"
    assert grid.scope == "scope2"
    assert grid.label_gu == "ગ્રીડ વીજળી"
    assert grid.label_hi == "ग्रिड बिजली"
    assert "bijli" in grid.synonyms

    # Verify EmissionFactor CEA versioning and dates
    ef_cea_recent = (
        db_session.query(EmissionFactor)
        .filter_by(activity_type="grid_electricity", source_version="v21.0")
        .first()
    )
    assert ef_cea_recent is not None
    assert float(ef_cea_recent.kgco2e_per_unit) == pytest.approx(0.710, 0.001)
    assert ef_cea_recent.reference_year == "FY2024-25"

    ef_cea_prev = (
        db_session.query(EmissionFactor)
        .filter_by(activity_type="grid_electricity", source_version="v20.0")
        .first()
    )
    assert ef_cea_prev is not None
    assert float(ef_cea_prev.kgco2e_per_unit) == pytest.approx(0.727, 0.001)
    assert ef_cea_prev.valid_to is not None

    # Verify Interventions
    interv = db_session.query(Intervention).filter_by(code="CA_LEAK_FIX").first()
    assert interv is not None
    assert interv.pool == "grid_kwh"
    assert interv.effect_type == "reduce_fraction"
    assert "share_of_pool" in interv.effect_params

    solar = db_session.query(Intervention).filter_by(code="SOLAR_ROOFTOP").first()
    assert solar is not None
    assert solar.effect_type == "onsite_generation"
    assert solar.circularity_points == 2

    # Verify Benchmarks
    bm = (
        db_session.query(Benchmark)
        .filter_by(industry="brass_components", metric="kwh_per_t_output")
        .first()
    )
    assert bm is not None
    assert float(bm.value_mode) == 2100.0


def test_standard_template_excel_structure():
    """Verify data/templates/decarbo_template.xlsx structure and data validation."""
    template_path = BASE_DIR / "data" / "templates" / "decarbo_template.xlsx"
    assert template_path.exists(), "decarbo_template.xlsx must exist"

    wb = openpyxl.load_workbook(str(template_path))
    assert "Template" in wb.sheetnames
    assert "How to fill" in wb.sheetnames

    ws = wb["Template"]
    expected_headers = [
        "month",
        "activity",
        "quantity",
        "unit",
        "cost_inr",
        "recycled_share",
        "distance_km",
        "vehicle",
        "notes",
    ]
    actual_headers = [cell.value for cell in ws[1][: len(expected_headers)]]
    assert actual_headers == expected_headers

    # Check sample row exists
    assert ws.max_row >= 2
    first_row_vals = [cell.value for cell in ws[2][: len(expected_headers)]]
    assert first_row_vals[0] == "2026-04"
    assert first_row_vals[1] == "grid_electricity"


def test_demo_factory_data_and_drift():
    """Verify demo_factory.xlsx has 12 months with festival dip and drift."""
    demo_xlsx = BASE_DIR / "data" / "demo" / "demo_factory.xlsx"
    assert demo_xlsx.exists(), "demo_factory.xlsx must exist"

    wb = openpyxl.load_workbook(str(demo_xlsx), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = rows[1:]

    month_idx = header.index("month")
    act_idx = header.index("activity")
    qty_idx = header.index("quantity")

    months = sorted(list({r[month_idx] for r in data if r[month_idx]}))
    assert len(months) == 12, "Demo data must span 12 months"

    # Monthly electricity totals and output totals
    monthly_kwh = {}
    monthly_output = {}
    for r in data:
        m = r[month_idx]
        if r[act_idx] == "grid_electricity":
            monthly_kwh[m] = monthly_kwh.get(m, 0.0) + float(r[qty_idx])
        elif r[act_idx] == "production_output":
            monthly_output[m] = monthly_output.get(m, 0.0) + float(r[qty_idx])

    # Festival dip in Oct/Nov (Diwali season)
    assert monthly_kwh["2025-11"] < monthly_kwh["2025-09"]
    assert monthly_kwh["2025-11"] < monthly_kwh["2025-12"]

    # Electricity intensity (kWh / tonne output) drift in last 3 months (+14%)
    baseline_months = ["2025-09", "2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05"]
    avg_baseline_kwh_per_t = sum(monthly_kwh[m] / monthly_output[m] for m in baseline_months) / len(
        baseline_months
    )
    drift_months = ["2026-06", "2026-07", "2026-08"]
    avg_drift_kwh_per_t = sum(monthly_kwh[m] / monthly_output[m] for m in drift_months) / len(
        drift_months
    )
    drift_pct = (avg_drift_kwh_per_t - avg_baseline_kwh_per_t) / avg_baseline_kwh_per_t

    assert drift_pct > 0.10, f"Expected drift > 10%, got {drift_pct * 100:.1f}%"


def test_tally_and_bill_files():
    """Verify tally purchase register CSV and sample electricity bill PDFs."""
    tally_csv = BASE_DIR / "data" / "demo" / "tally_purchase_register.csv"
    assert tally_csv.exists(), "tally_purchase_register.csv must exist"

    with open(tally_csv, encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader)
        expected = [
            "Date",
            "Particulars",
            "Voucher Type",
            "Voucher No",
            "Quantity",
            "Unit",
            "Rate",
            "Gross Value",
        ]
        assert header == expected
        rows = list(reader)
        assert len(rows) >= 12, "Tally register should have at least 12 rows"

    bills_dir = BASE_DIR / "data" / "demo" / "bills"
    assert bills_dir.exists(), "bills directory must exist"
    bill_files = list(bills_dir.glob("*.pdf"))
    assert len(bill_files) == 3, f"Expected 3 bill PDFs, found {len(bill_files)}"

    for pdf_path in bill_files:
        assert pdf_path.stat().st_size > 500, f"{pdf_path.name} is too small"
        with open(pdf_path, "rb") as f:
            header = f.read(5)
            assert header.startswith(b"%PDF-"), f"{pdf_path.name} is not a valid PDF"
