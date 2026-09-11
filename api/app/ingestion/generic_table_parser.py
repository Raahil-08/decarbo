"""Generic table (Tally exports / custom spreadsheets) parser and mapper per PRD §9.4."""

import csv
from io import BytesIO
from typing import Any

import openpyxl

from app.engine.units import convert_to_canonical, normalize_unit_string
from app.ingestion.llm import LLMProvider, TableMappingResult, get_llm_provider


def extract_table_structure(
    file_bytes: bytes,
    filename: str,
) -> tuple[str, list[list[Any]], list[list[Any]], list[str]]:
    """Extract sheet name, candidate header rows, sample rows, and distinct item labels."""
    raw_rows: list[list[Any]] = []
    sheet_name = "Sheet1"

    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(text.splitlines())
        raw_rows = [list(r) for r in reader if any(r)]
    else:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            if any(row):
                raw_rows.append(list(row))

    if not raw_rows:
        return sheet_name, [], [], []

    # First 15 rows as header/structure candidates
    header_candidates = raw_rows[:15]
    sample_rows = raw_rows[1:16]

    # Find item column from candidate headers
    item_col_idx = None
    item_col_names = ["particulars", "item name", "item", "description", "product", "material"]
    for row in header_candidates:
        for c_idx, cell in enumerate(row):
            if cell and str(cell).strip().lower() in item_col_names:
                item_col_idx = c_idx
                break
        if item_col_idx is not None:
            break

    item_candidates = set()
    for row in raw_rows[1:]:
        if item_col_idx is not None and item_col_idx < len(row):
            val = row[item_col_idx]
            if val is not None:
                s = str(val).strip()
                if len(s) > 1:
                    item_candidates.add(s)
        else:
            for val in row:
                if val is not None:
                    s = str(val).strip()
                    if len(s) > 2 and not s.replace(".", "").replace("-", "").isdigit():
                        item_candidates.add(s)

    distinct_items = sorted(list(item_candidates))[:150]
    return sheet_name, header_candidates, sample_rows, distinct_items



def propose_generic_table_mapping(
    file_bytes: bytes,
    filename: str,
    canonical_activities: list[dict[str, Any]],
    provider: LLMProvider | None = None,
) -> tuple[TableMappingResult, list[list[Any]]]:
    """Propose column and item mapping for generic table using LLM or rule-based fallback."""
    if provider is None:
        provider = get_llm_provider()

    sheet_name, header_candidates, sample_rows, distinct_items = extract_table_structure(
        file_bytes, filename
    )

    mapping_res = provider.propose_table_mapping(
        sheet_name=sheet_name,
        header_candidates=header_candidates,
        sample_rows=sample_rows,
        distinct_items=distinct_items,
        canonical_activities=canonical_activities,
    )

    # Return full rows as well
    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8", errors="replace")
        all_rows = list(csv.reader(text.splitlines()))
    else:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        ws = wb[mapping_res.sheet] if mapping_res.sheet in wb.sheetnames else wb.active
        all_rows = list(ws.iter_rows(values_only=True))

    return mapping_res, all_rows


def apply_confirmed_table_mapping(
    all_rows: list[list[Any]],
    mapping: TableMappingResult,
    activity_types: list[Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Transform rows into activity records using confirmed column and item mappings."""
    if not all_rows or len(all_rows) <= mapping.header_row_index:
        return [], [{"row": 0, "field": "mapping", "code": "NO_DATA", "message_key": "errors.no_data", "severity": "block"}]

    header = [str(c).strip().lower() if c is not None else "" for c in all_rows[mapping.header_row_index]]
    cols = {k: v.lower() for k, v in mapping.columns.items()}

    # Resolve column indexes
    def find_idx(target_name):
        t = cols.get(target_name, "").lower()
        if not t:
            return None
        for idx, h in enumerate(header):
            if h == t or t in h:
                return idx
        return None

    date_idx = find_idx("date_or_month")
    item_idx = find_idx("item")
    qty_idx = find_idx("quantity")
    unit_idx = find_idx("unit")
    cost_idx = find_idx("cost_inr")

    if item_idx is None or qty_idx is None:
        return [], [{"row": 0, "field": "columns", "code": "MISSING_MAPPING", "message_key": "errors.missing_column_mapping", "severity": "block"}]

    # Build item classification map
    item_to_act: dict[str, str | None] = {}
    for im in mapping.item_mappings:
        item_to_act[im.source_label.strip().lower()] = im.activity_type

    act_type_map = {getattr(a, "key", ""): a for a in activity_types}
    draft_records = []
    issues = []

    for row_num, row in enumerate(all_rows[mapping.header_row_index + 1:], start=mapping.header_row_index + 2):
        if not any(row):
            continue

        raw_item = row[item_idx] if item_idx < len(row) else None
        if not raw_item:
            continue

        item_str = str(raw_item).strip()
        act_key = item_to_act.get(item_str.lower())
        if not act_key:
            # Item mapped to None (non-emission or skipped)
            continue

        act_meta = act_type_map.get(act_key)
        if not act_meta:
            continue

        # Extract quantity
        raw_qty = row[qty_idx] if qty_idx < len(row) else None
        try:
            qty_float = float(raw_qty)
            if qty_float <= 0:
                continue
        except (ValueError, TypeError):
            continue

        # Extract unit
        raw_u = row[unit_idx] if unit_idx is not None and unit_idx < len(row) else getattr(act_meta, "canonical_unit", "")
        unit_str = str(raw_u).strip() if raw_u else getattr(act_meta, "canonical_unit", "")
        canon_unit = getattr(act_meta, "canonical_unit", unit_str)

        try:
            norm_u = normalize_unit_string(unit_str)
            qty_canon = convert_to_canonical(qty_float, norm_u, canon_unit)
        except Exception:
            qty_canon = qty_float

        # Extract Date / Month
        raw_d = row[date_idx] if date_idx is not None and date_idx < len(row) else None
        d_str = str(raw_d).strip() if raw_d else "2026-04-01"
        month_str = d_str[:7].replace("/", "-")
        if len(month_str) < 7:
            month_str = "2026-04"

        # Extract Cost
        raw_c = row[cost_idx] if cost_idx is not None and cost_idx < len(row) else None
        cost_inr = None
        if raw_c:
            try:
                cost_inr = float(raw_c)
            except ValueError:
                pass

        draft_records.append({
            "period_month": f"{month_str}-01",
            "activity_type": act_key,
            "quantity": qty_float,
            "unit": unit_str,
            "quantity_canonical": qty_canon,
            "cost_inr": cost_inr,
            "source_note": f"row {row_num}: {item_str}",
            "confidence": 0.90,
            "attributes": {"source_item": item_str},
        })

    return draft_records, issues
