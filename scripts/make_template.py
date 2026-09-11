"""Generate standard upload template decarbo_template.xlsx with data validation."""

import csv
from pathlib import Path
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_DIR = BASE_DIR / "data" / "seed"
TEMPLATES_DIR = BASE_DIR / "data" / "templates"


def load_activity_keys():
    csv_file = SEED_DIR / "activity_types.csv"
    keys = []
    if csv_file.exists():
        with open(csv_file, mode="r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                keys.append(row["key"].strip())
    return keys


def generate_template():
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    out_path = TEMPLATES_DIR / "decarbo_template.xlsx"

    wb = openpyxl.Workbook()
    # Sheet 1: Template
    ws_template = wb.active
    ws_template.title = "Template"

    # Sheet 2: How to fill
    ws_guide = wb.create_sheet(title="How to fill")

    # Sheet 3: Hidden reference for validation list
    ws_ref = wb.create_sheet(title="_Reference")
    ws_ref.sheet_state = "hidden"

    activity_keys = load_activity_keys()
    for idx, key in enumerate(activity_keys, start=1):
        ws_ref.cell(row=idx, column=1, value=key)

    # Styles
    header_fill = PatternFill(start_color="1D2A45", end_color="1D2A45", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(
        left=Side(style="thin", color="D6DAE1"),
        right=Side(style="thin", color="D6DAE1"),
        top=Side(style="thin", color="D6DAE1"),
        bottom=Side(style="thin", color="D6DAE1"),
    )

    headers = [
        ("month", "YYYY-MM (e.g. 2026-04)", align_center, 14),
        ("activity", "Activity Key from dropdown", align_left, 28),
        ("quantity", "Numeric value", align_right, 14),
        ("unit", "Unit (kWh, L, kg, t, etc.)", align_center, 12),
        ("cost_inr", "Optional spend in INR", align_right, 16),
        ("recycled_share", "Materials only (0.0 to 1.0)", align_right, 18),
        ("distance_km", "Transport only (km)", align_right, 16),
        ("vehicle", "Transport only (hgv / lcv)", align_center, 16),
        ("notes", "Optional note / reference", align_left, 30),
    ]

    for col_idx, (col_name, _, align, width) in enumerate(headers, start=1):
        cell = ws_template.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align
        cell.border = border_thin
        ws_template.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # Sample rows (PRD §9.2)
    sample_rows = [
        ("2026-04", "grid_electricity", 95000, "kWh", 812000, "", "", "", "Main factory electricity bill"),
        ("2026-04", "diesel", 1500, "L", 135000, "", "", "", "Backup DG set operations"),
        ("2026-04", "brass_input_primary", 45000, "kg", 18000000, 0.2, "", "", "IS319 extruded brass rod"),
    ]

    for row_idx, row_data in enumerate(sample_rows, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws_template.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            cell.alignment = headers[col_idx - 1][2]
            if isinstance(val, (int, float)):
                cell.font = Font(name="Arial", size=10)

    # Data validation for activity column (rows 2 to 1000)
    dv = DataValidation(
        type="list",
        formula1=f"='_Reference'!$A$1:$A${len(activity_keys)}",
        allow_blank=True,
    )
    dv.error = "Please choose a valid activity from the dropdown list"
    dv.errorTitle = "Invalid Activity"
    ws_template.add_data_validation(dv)
    dv.add("B2:B1000")

    # Sheet 2: Guide & Instructions
    ws_guide.column_dimensions["A"].width = 20
    ws_guide.column_dimensions["B"].width = 35
    ws_guide.column_dimensions["C"].width = 45

    guide_headers = ["Column / કૉલમ", "English Instruction", "ગુજરાતી માર્ગદર્શિકા (Gujarati Guide)"]
    for c_idx, gh in enumerate(guide_headers, start=1):
        c = ws_guide.cell(row=1, column=c_idx, value=gh)
        c.fill = header_fill
        c.font = header_font
        c.alignment = align_left
        c.border = border_thin

    guide_rows = [
        ("month", "Month in YYYY-MM format (e.g. 2026-04).", "વર્ષ અને મહિનો YYYY-MM ફોર્મેટમાં (દા.ત. 2026-04)."),
        ("activity", "Canonical activity key. Use the dropdown list.", "પ્રવૃત્તિનું નામ. ડ્રોપડાઉન યાદીમાંથી પસંદ કરો."),
        ("quantity", "Numeric activity quantity without commas.", "સંખ્યાત્મક મૂલ્ય અલ્પવિરામ વગર."),
        ("unit", "Unit of measurement (kWh, L, kg, t, m3).", "માપન એકમ (kWh, L, kg, t, m3)."),
        ("cost_inr", "Optional amount spent in Rupees (INR).", "વૈકલ્પિક ખર્ચ રૂપિયામાં (INR)."),
        ("recycled_share", "For metal/material rows: fraction recycled (0.0 to 1.0).", "ધાતુ/સામગ્રી માટે: રિસાઇકલ્ડ હિસ્સો (0.0 થી 1.0)."),
        ("distance_km", "For transport rows: distance travelled in km.", "પરિવહન માટે: કાપેલું અંતર કિમીમાં."),
        ("vehicle", "For transport: 'hgv' (heavy truck) or 'lcv' (tempo/pickup).", "પરિવહન માટે: 'hgv' (મોટો ટ્રક) અથવા 'lcv' (ટેમ્પો)."),
        ("notes", "Any reference, invoice number, or meter details.", "કોઈપણ સંદર્ભ, ઇન્વોઇસ નંબર અથવા મીટર વિગતો."),
    ]

    for r_idx, g_row in enumerate(guide_rows, start=2):
        for c_idx, val in enumerate(g_row, start=1):
            c = ws_guide.cell(row=r_idx, column=c_idx, value=val)
            c.border = border_thin
            c.alignment = align_left

    wb.save(out_path)
    print(f"Generated upload template at: {out_path}")
    return out_path


if __name__ == "__main__":
    generate_template()
